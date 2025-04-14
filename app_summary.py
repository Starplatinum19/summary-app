import streamlit as st
import pdfplumber
import re
from langdetect import detect
from transformers import pipeline, MarianMTModel, MarianTokenizer
from sumy.parsers.plaintext import PlaintextParser
from sumy.nlp.tokenizers import Tokenizer
from sumy.summarizers.text_rank import TextRankSummarizer
from docx import Document

def clean_text(text):
    text = re.sub(r'\[\d+\]', '', text)
    text = re.sub(r'\([A-Za-z, ]+\d{4}\)', '', text)
    text = re.sub(r'[^a-zA-Z0-9áéíóúñÁÉÍÓÚÑ\s\.,;:¡!¿?\-]', '', text)
    text = text.lower()
    text = re.sub(r'\n\s*\n', '\n', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text

st.title("📄 Automatic Research Paper Summarizer")

uploaded_pdf = st.file_uploader("1️⃣ Upload your PDF file", type=["pdf"])

if uploaded_pdf:
    with pdfplumber.open(uploaded_pdf) as pdf:
        raw_text = ""
        for page in pdf.pages:
            raw_text += page.extract_text() + "\n"

    cleaned_text = clean_text(raw_text)
    lang = detect(cleaned_text)
    st.write(f"🌍 Detected Language: {lang}")

    translate = st.radio("2️⃣ Translate text to Spanish?", ["No", "Yes (if in English)"])
    if translate == "Yes (if in English)" and lang == "en":
        st.info("Translating to Spanish...")
        model_id = "Helsinki-NLP/opus-mt-en-es"
        tokenizer = MarianTokenizer.from_pretrained(model_id)
        model = MarianMTModel.from_pretrained(model_id)
        tokens = tokenizer(cleaned_text[:1000], return_tensors="pt", padding=True, truncation=True)
        translation = model.generate(**tokens)
        cleaned_text = tokenizer.decode(translation[0], skip_special_tokens=True)
        st.success("✅ Translated to Spanish")

    summary_type = st.radio("3️⃣ Choose summary type", ["Extractive", "Abstractive"])
    summary = ""

    if st.button("📝 Generate Summary"):
        if summary_type == "Extractive":
            parser = PlaintextParser.from_string(cleaned_text, Tokenizer("spanish"))
            summarizer = TextRankSummarizer()
            summary = "\n".join(str(sentence) for sentence in summarizer(parser.document, 5))
        else:
            summarizer = pipeline("summarization", model="facebook/bart-large-cnn")
            summary = summarizer(cleaned_text[:1024], max_length=130, min_length=30, do_sample=False)[0]["summary_text"]

        st.subheader("🧾 Summary Output")
        st.write(summary)

        st.download_button("📥 Download as TXT", summary, file_name="summary.txt")

        doc = Document()
        doc.add_heading("Generated Summary", level=1)
        doc.add_paragraph(summary)
        doc.save("summary.docx")
        with open("summary.docx", "rb") as f:
            st.download_button("📥 Download as DOCX", f, file_name="summary.docx")
